"""Exportacion a Excel (.xlsx) con formato de reporte (CU11 paso 6, CU28).

Hoja "Reporte": encabezado (empresa, titulo, filtros aplicados, fecha de
generacion y usuario), la tabla con formato y una fila de totales.
Hoja "Indicadores": las secciones de resumen que se le pasen.
Hoja "Datos": los mismos registros sin formato, para analizarlos.
"""

import io
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

_TITULO = Font(bold=True, size=14, color="1F2937")
_NEGRITA = Font(bold=True)
_CABECERA = Font(bold=True, color="FFFFFF")
_RELLENO = PatternFill("solid", fgColor="4F46E5")
_LINEA = Border(bottom=Side(style="thin", color="D1D5DB"))
FORMATO_BS = '#,##0.00 "Bs"'


def libro_excel(
    *,
    titulo: str,
    empresa: str,
    usuario: str,
    filtros: list[tuple[str, str]],
    encabezados: list[str],
    filas: list[list],
    columnas_monto: tuple[int, ...] = (),
    indicadores: list[tuple[str, list[str], list[list]]] = (),
    generado: datetime | None = None,
) -> bytes:
    libro = Workbook()
    hoja = libro.active
    hoja.title = "Reporte"
    hoja["A1"] = empresa
    hoja["A1"].font = _TITULO
    hoja["A2"] = titulo
    hoja["A2"].font = _NEGRITA
    hoja["A3"] = f"Generado el {(generado or datetime.now()).strftime('%d/%m/%Y %H:%M')} por {usuario}"
    fila = 4
    for nombre, valor in filtros:
        hoja.cell(row=fila, column=1, value=nombre).font = _NEGRITA
        hoja.cell(row=fila, column=2, value=valor)
        fila += 1
    fila += 1

    inicio = fila
    for col, texto in enumerate(encabezados, start=1):
        celda = hoja.cell(row=fila, column=col, value=texto)
        celda.font, celda.fill = _CABECERA, _RELLENO
        celda.alignment = Alignment(horizontal="center", vertical="center")
    for valores in filas:
        fila += 1
        for col, valor in enumerate(valores, start=1):
            celda = hoja.cell(row=fila, column=col, value=valor)
            celda.border = _LINEA
            if col - 1 in columnas_monto:
                celda.number_format = FORMATO_BS
            elif isinstance(valor, datetime):
                celda.number_format = "dd/mm/yyyy hh:mm"
    if filas:
        fila += 1
        hoja.cell(row=fila, column=1, value=f"Total: {len(filas)} registro(s)").font = _NEGRITA
        for col in columnas_monto:
            letra = get_column_letter(col + 1)
            celda = hoja.cell(row=fila, column=col + 1, value=f"=SUM({letra}{inicio + 1}:{letra}{fila - 1})")
            celda.font, celda.number_format = _NEGRITA, FORMATO_BS
        hoja.auto_filter.ref = f"A{inicio}:{get_column_letter(len(encabezados))}{fila - 1}"
    hoja.freeze_panes = hoja.cell(row=inicio + 1, column=1)
    for col, texto in enumerate(encabezados, start=1):
        ancho = max([len(str(texto))] + [len(str(f[col - 1])) for f in filas[:300] if f[col - 1] is not None])
        hoja.column_dimensions[get_column_letter(col)].width = min(max(ancho + 2, 10), 45)

    if indicadores:
        resumen = libro.create_sheet("Indicadores")
        fila = 1
        for seccion, cabecera, datos in indicadores:
            resumen.cell(row=fila, column=1, value=seccion).font = _TITULO
            fila += 1
            for col, texto in enumerate(cabecera, start=1):
                celda = resumen.cell(row=fila, column=col, value=texto)
                celda.font, celda.fill = _CABECERA, _RELLENO
            for valores in datos:
                fila += 1
                for col, valor in enumerate(valores, start=1):
                    resumen.cell(row=fila, column=col, value=valor)
            fila += 2
        for col in range(1, 6):
            resumen.column_dimensions[get_column_letter(col)].width = 28

    crudos = libro.create_sheet("Datos")
    crudos.append(encabezados)
    for valores in filas:
        crudos.append(valores)

    salida = io.BytesIO()
    libro.save(salida)
    return salida.getvalue()

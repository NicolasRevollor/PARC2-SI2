"""CU11: alcance por sucursal (flujo alternativo 2) y exportacion a Excel (paso 6)."""

import io
from datetime import datetime

from openpyxl import load_workbook

from app.services.alcance import recortar, sucursales_visibles
from app.services.excel import libro_excel


def test_el_administrador_ve_todas_las_sucursales():
    assert sucursales_visibles(["tenant.admin", "devolucion.ver"], [1]) is None


def test_el_encargado_asignado_solo_ve_sus_sucursales():
    assert sucursales_visibles(["devolucion.ver"], [1, 4]) == {1, 4}


def test_el_personal_sin_asignar_ve_toda_la_empresa():
    assert sucursales_visibles(["devolucion.ver"], []) is None


def test_el_filtro_de_sucursal_no_amplia_el_alcance():
    assert recortar(None, {1}) == {1}
    assert recortar(1, {1, 2}) == {1}
    assert recortar(3, {1, 2}) == set()
    assert recortar(3, None) == {3}
    assert recortar(None, None) is None


def test_excel_con_reporte_indicadores_y_datos_sin_formato():
    contenido = libro_excel(
        titulo="Historial de devoluciones",
        empresa="Maison Oriente",
        usuario="Roberto Encargado",
        filtros=[("Sucursales", "Maison Oriente Norte"), ("Motivo", "Defecto de fabrica")],
        encabezados=["Codigo", "Fecha", "Reembolsado"],
        filas=[["DEV-000001", datetime(2026, 10, 20, 15, 30), 120.0], ["DEV-000002", datetime(2026, 10, 21, 9, 0), 80.5]],
        columnas_monto=(2,),
        indicadores=[("Motivos", ["Motivo", "Casos"], [["Defecto de fabrica", 2]])],
    )
    libro = load_workbook(io.BytesIO(contenido))
    assert libro.sheetnames == ["Reporte", "Indicadores", "Datos"]
    reporte = libro["Reporte"]
    assert reporte["A1"].value == "Maison Oriente"
    valores = [c.value for fila in reporte.iter_rows() for c in fila if c.value is not None]
    assert "Defecto de fabrica" in valores and "DEV-000002" in valores
    assert any(isinstance(v, str) and v.startswith("=SUM(") for v in valores)
    datos = list(libro["Datos"].iter_rows(values_only=True))
    assert datos[0] == ("Codigo", "Fecha", "Reembolsado")
    assert datos[2][2] == 80.5 and len(datos) == 3
